"""
fs_gap_engine.py  -  Find out WHY a balance sheet does not tally.

Works on any Excel financial-statement workbook that is built with formulas
(SUMIF / SUMIFS links to a trial balance, direct cell links, or a mix).

Pipeline
  1. Locate the balance sheet (Total Assets vs Total Equity & Liabilities) and its period columns.
  2. For every period column that does not tally, trace each statement line back through the
     note sheets to the trial-balance (TB) cells it reads.
  3. Run checks:  internal arithmetic, TB rows used twice, TB rows used nowhere, partly linked rows,
     same line pulled from two period columns, #REF!/errors, formulas that differ from the
     neighbouring period columns, hard-coded numbers, wrong-sign balances.
  4. Search for the smallest set of findings whose combined effect equals the gap.
  5. Write the result as a readable analysis (markdown) and an Excel report.

Nothing is changed in the user's workbook.
"""
from __future__ import annotations

import bisect
import datetime as dt
import io
import itertools
import math
import re
import warnings
from collections import Counter, defaultdict

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

warnings.filterwarnings("ignore")

ERRORS = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}
MAX_RANGE_CELLS = 50000


# --------------------------------------------------------------------------- small helpers
def is_err(v):
    return isinstance(v, str) and v.strip() in ERRORS


def num(v):
    if isinstance(v, bool):
        return None
    return float(v) if isinstance(v, (int, float)) else None


def is_formula(v):
    return isinstance(v, str) and v.startswith("=") and len(v) > 1


def addr(col, row):
    return f"{get_column_letter(col)}{row}"


def split_addr(a):
    c, r = coordinate_from_string(a)
    return column_index_from_string(c), r


def fmt(v, dec=2):
    if v is None:
        return "-"
    if isinstance(v, str):
        return v
    s = f"{abs(v):,.{dec}f}"
    return f"-{s}" if v < -0.5 * 10 ** (-dec) else s


def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


# --------------------------------------------------------------------------- formula parsing
class Operand:
    __slots__ = ("kind", "value")

    def __init__(self, kind, value):
        self.kind, self.value = kind, value


class Func:
    __slots__ = ("name", "args")

    def __init__(self, name, args):
        self.name, self.args = name, args


class Group:
    __slots__ = ("items",)

    def __init__(self, items):
        self.items = items


def parse_formula(f):
    try:
        tk = Tokenizer(f)
    except Exception:
        return None
    root = {"args": [[]]}
    stack = [root]
    for t in tk.items:
        if t.type == Token.FUNC:
            if t.subtype == Token.OPEN:
                stack.append({"name": t.value[:-1].upper(), "args": [[]], "kind": "f"})
            elif len(stack) > 1:
                fr = stack.pop()
                stack[-1]["args"][-1].append(Func(fr["name"], fr["args"]))
        elif t.type == Token.PAREN:
            if t.subtype == Token.OPEN:
                stack.append({"args": [[]], "kind": "g"})
            elif len(stack) > 1:
                fr = stack.pop()
                stack[-1]["args"][-1].append(Group(fr["args"][0]))
        elif t.type == Token.SEP:
            stack[-1]["args"].append([])
        elif t.type == Token.OPERAND:
            stack[-1]["args"][-1].append(Operand(t.subtype, t.value))
        elif t.type in (Token.OP_IN, Token.OP_PRE):
            stack[-1]["args"][-1].append(("op", t.value, t.type == Token.OP_PRE))
        elif t.type == Token.OP_POST:
            stack[-1]["args"][-1].append(("op", "%", False))
    while len(stack) > 1:  # unbalanced formula
        fr = stack.pop()
        stack[-1]["args"][-1].extend(fr["args"][0])
    return root["args"][0]


def collect_operands(items, out):
    for it in items:
        if isinstance(it, Operand):
            out.append(it)
        elif isinstance(it, Group):
            collect_operands(it.items, out)
        elif isinstance(it, Func):
            for a in it.args:
                collect_operands(a, out)


def crit_match(v, crit):
    """Excel-style SUMIF criteria match."""
    if v is None:
        return crit in ("", "=")
    if isinstance(crit, (int, float)) and not isinstance(crit, bool):
        if isinstance(v, (int, float)):
            return math.isclose(v, crit, rel_tol=0, abs_tol=1e-9)
        try:
            return float(str(v).strip()) == crit
        except ValueError:
            return False
    s = str(crit)
    m = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", s, re.S)
    op, rhs = (m.group(1) or "="), m.group(2)
    try:
        rf = float(rhs)
        try:
            vf = float(v) if not isinstance(v, str) else float(v.strip())
        except ValueError:
            return op == "<>"
        return {"=": vf == rf, "<>": vf != rf, "<": vf < rf, ">": vf > rf, "<=": vf <= rf, ">=": vf >= rf}[op]
    except ValueError:
        pass
    vs = str(v)
    if "*" in rhs or "?" in rhs:
        pat = re.compile("^" + re.escape(rhs).replace(r"\*", ".*").replace(r"\?", ".") + "$", re.I)
        hit = bool(pat.match(vs))
    else:
        hit = vs.lower() == rhs.lower()
    return hit if op == "=" else (not hit if op == "<>" else False)


_NOEVAL = object()


# --------------------------------------------------------------------------- workbook wrapper
class Book:
    def __init__(self, data: bytes, name="workbook.xlsx"):
        self.name = name
        self.wf = load_workbook(io.BytesIO(data))                  # formulas
        self.wv = load_workbook(io.BytesIO(data), data_only=True)  # cached values
        self.sheetnames = self.wf.sheetnames
        self.warnings: list[str] = []
        self._exp: dict = {}
        self._mult: dict = {}
        self._scl: dict = {}
        self._ctx: dict = {}

    def warn(self, msg):
        if msg not in self.warnings and len(self.warnings) < 60:
            self.warnings.append(msg)

    def raw(self, sh, a):
        return self.wf[sh][a].value

    def val(self, sh, a):
        return self.wv[sh][a].value

    def maxrow(self, sh):
        return min(self.wf[sh].max_row or 1, 100000)

    def maxcol(self, sh):
        return min(self.wf[sh].max_column or 1, 200)

    def rect_values(self, sh, c1, r1, c2, r2):
        return [list(r) for r in self.wv[sh].iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True)]

    # ---- reference resolution
    def resolve_rect(self, ref, home):
        if "!" in ref:
            sh, rg = ref.rsplit("!", 1)
            sh = sh.strip()
            if sh.startswith("'") and sh.endswith("'"):
                sh = sh[1:-1].replace("''", "'")
        else:
            sh, rg = home, ref
        if sh not in self.wf.sheetnames:
            self.warn(f"Reference to a sheet that is not in the workbook (external link?): {ref}")
            return None
        try:
            c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        except Exception:
            self.warn(f"Could not resolve reference '{ref}' (named range?)")
            return None
        mr, mc = self.maxrow(sh), self.maxcol(sh)
        r1 = 1 if r1 is None else r1
        r2 = mr if r2 is None else r2
        c1 = 1 if c1 is None else c1
        c2 = mc if c2 is None else c2
        r2 = min(r2, max(mr, r1))
        c2 = min(c2, max(mc, c1))
        return sh, c1, r1, c2, r2

    def resolve_cells(self, ref, home):
        rect = self.resolve_rect(ref, home)
        if not rect:
            return []
        sh, c1, r1, c2, r2 = rect
        if (c2 - c1 + 1) * (r2 - r1 + 1) > MAX_RANGE_CELLS:
            self.warn(f"Range too large to trace, skipped: {ref}")
            return []
        return [(sh, addr(c, r)) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    # ---- one-step dependency expansion of a formula cell (memoised)
    def expand(self, sh, a):
        key = (sh, a)
        if key in self._exp:
            return self._exp[key]
        raw = self.raw(sh, a)
        out = []
        items = parse_formula(raw) if is_formula(raw) else None
        if items is not None:
            self._walk(items, 1, sh, out, False)
        res, seen, mult, scl = [], set(), Counter(), {}
        for (csh, ca), sg, sc in out:
            k = (csh, ca, sg)
            mult[k] += 1
            scl[k] = scl.get(k, False) or sc
            if k not in seen:
                seen.add(k)
                res.append(k)
        self._exp[key] = res
        self._mult[key] = mult
        self._scl[key] = scl
        return res

    @staticmethod
    def _term_scaled(items):
        """True when a term multiplies/divides by something that is not a plain unit conversion (rate, % or another cell)."""
        ops = [it[1] for it in items if isinstance(it, tuple)]
        if "%" in ops:
            return True
        if not any(o in ("*", "/", "^") for o in ops):
            return False
        n_ref, nums, prev = 0, [], None
        for it in items:
            if isinstance(it, tuple):
                prev = it[1]
                continue
            if isinstance(it, Operand) and it.kind == "NUMBER":
                if prev != "^":
                    try:
                        nums.append(float(it.value))
                    except ValueError:
                        pass
            else:
                n_ref += 1
            prev = None
        if n_ref >= 2:
            return True
        unit = {1.0, 10.0, 100.0, 1000.0, 10000.0, 100000.0, 1000000.0, 10000000.0, 12.0}
        return any(n not in unit for n in nums)

    @staticmethod
    def _term_factor(items):
        """numeric factor of plain unit conversions in a term, e.g. /1000 or /10^5 (refs count as 1)"""
        nf, prev = [], None
        for it in items:
            if isinstance(it, tuple):
                prev = it[1]
                continue
            if isinstance(it, Operand) and it.kind == "NUMBER":
                try:
                    v = float(it.value)
                except ValueError:
                    v = None
                if v is not None:
                    if prev == "^" and nf:
                        nf[-1][1] = nf[-1][1] ** v
                    else:
                        nf.append([prev if prev in ("*", "/") else "*", v])
            prev = None
        f = 1.0
        for op, v in nf:
            f = f * v if op == "*" else f / (v or 1.0)
        return f

    def _walk(self, items, sign, home, out, scaled=False):
        terms, cur, cur_sign = [], [], sign
        for it in items:
            if isinstance(it, tuple) and it[1] in ("+", "-"):
                if it[2]:                                   # prefix operator
                    if not cur:
                        cur_sign = -cur_sign if it[1] == "-" else cur_sign
                    else:
                        cur.append(it)
                else:                                       # additive infix: start a new term
                    terms.append((cur, cur_sign))
                    cur, cur_sign = [], (sign if it[1] == "+" else -sign)
            else:
                cur.append(it)
        terms.append((cur, cur_sign))
        for titems, tsign in terms:
            sc = scaled or self._term_scaled(titems)
            if not sc:
                tsign = tsign * self._term_factor(titems)
            for it in titems:
                if isinstance(it, Operand):
                    if it.kind == "RANGE":
                        for c in self.resolve_cells(it.value, home):
                            out.append((c, tsign, sc))
                elif isinstance(it, Group):
                    self._walk(it.items, tsign, home, out, sc)
                elif isinstance(it, Func):
                    self._func(it, tsign, home, out, sc)

    def _func(self, fn, sign, home, out, scaled=False):
        name = fn.name.split(".")[-1]
        if name in ("SUMIF", "SUMIFS"):
            if self._sumif(fn, name == "SUMIFS", sign, home, out, scaled):
                return
            self.warn("A SUMIF/SUMIFS criteria could not be evaluated; its whole sum range was treated as linked")
        args = fn.args[1:] if name == "IF" else fn.args
        for a in args:
            self._walk(a, sign, home, out, scaled)

    def _eval_crit(self, arg, home):
        its = [x for x in arg if not isinstance(x, tuple)]
        ops = [x for x in arg if isinstance(x, tuple)]
        parts = []
        for it in its:
            if not isinstance(it, Operand):
                return _NOEVAL
            if it.kind == "RANGE":
                rect = self.resolve_rect(it.value, home)
                if not rect or rect[1] != rect[3] or rect[2] != rect[4]:
                    return _NOEVAL
                parts.append(self.val(rect[0], addr(rect[1], rect[2])))
            elif it.kind == "NUMBER":
                parts.append(float(it.value))
            elif it.kind == "TEXT":
                parts.append(it.value[1:-1].replace('""', '"'))
            else:
                return _NOEVAL
        if len(parts) == 1:
            return 0.0 if parts[0] is None else parts[0]
        if parts and len(ops) == len(parts) - 1 and all(o[1] == "&" for o in ops):
            return "".join("" if p is None else (str(int(p)) if isinstance(p, float) and p == int(p) else str(p)) for p in parts)
        return _NOEVAL

    def _sumif(self, fn, ifs, sign, home, out, scaled=False):
        A = fn.args

        def single(arg):
            its = [x for x in arg if not isinstance(x, tuple)]
            return its[0] if (len(its) == 1 and isinstance(its[0], Operand) and its[0].kind == "RANGE") else None

        try:
            if not ifs:
                rng = single(A[0])
                sr = single(A[2]) if len(A) > 2 else rng
                pairs = [(rng, A[1])]
            else:
                sr = single(A[0])
                pairs = [(single(A[i]), A[i + 1]) for i in range(1, len(A) - 1, 2)]
        except IndexError:
            return False
        if sr is None or any(p[0] is None for p in pairs):
            return False
        srect = self.resolve_rect(sr.value, home)
        rects = [self.resolve_rect(p[0].value, home) for p in pairs]
        if not srect or any(r is None for r in rects):
            return False
        crits = [self._eval_crit(p[1], home) for p in pairs]
        if any(c is _NOEVAL for c in crits):
            return False
        vals = [self.rect_values(*r) for r in rects]
        H = min(r[4] - r[2] + 1 for r in rects)
        W = min(r[3] - r[1] + 1 for r in rects)
        ssh, sc1, sr1 = srect[0], srect[1], srect[2]
        for i in range(H):
            for j in range(W):
                if all(crit_match(v[i][j] if i < len(v) and j < len(v[i]) else None, c) for v, c in zip(vals, crits)):
                    out.append(((ssh, addr(sc1 + j, sr1 + i)), sign, scaled))
        return True


# --------------------------------------------------------------------------- locating the statement
RE_ASSETS = re.compile(r"^\W*assets?\W*$", re.I)
RE_EL = re.compile(r"^\W*(equity\s*(and|&)\s*liabilit(y|ies)|liabilit(y|ies)\s*(and|&)\s*equity)\W*$", re.I)
RE_TOTAL = re.compile(r"^\W*total\W*$", re.I)
RE_TOTAL_A = re.compile(r"^\W*total\s+assets\W*$", re.I)
RE_TOTAL_EL = re.compile(r"^\W*total\s+(equity\s*(and|&)\s*liabilit(y|ies)|liabilit(y|ies)\s*(and|&)\s*equity)\W*$", re.I)
RE_PERIOD = re.compile(
    r"(as\s+(at|on)|period|quarter|year|half|ended|\b\d{1,2}\s*[./\-]\s*\d{1,2}\s*[./\-]\s*\d{2,4}\b|\b(19|20)\d{2}\b|"
    r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\b)", re.I)
RE_DATE = re.compile(r"(\d{1,2})\s*[./\-]\s*(\d{1,2})\s*[./\-]\s*(\d{2,4})")
RE_LABEL_JUNK = re.compile(r"^\(?[a-zA-Z0-9]{1,3}\)?\.?$")


def parse_date(text):
    ms = RE_DATE.findall(str(text))
    if not ms:
        return None
    d, m, y = ms[-1]
    y = int(y)
    y += 2000 if y < 100 else 0
    try:
        return dt.date(y, int(m), int(d))
    except ValueError:
        return None


def sheet_texts(bk, sh, max_rows=400, max_cols=8):
    ws = bk.wv[sh]
    out = {}
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row or 1), max_col=max_cols, values_only=True), 1):
        for c, v in enumerate(row, 1):
            if isinstance(v, str) and v.strip():
                out.setdefault(r, []).append((c, v.strip()))
    return out


def find_structure(bk, sh):
    tx = sheet_texts(bk, sh)
    a_hdr = l_hdr = None
    tot_a, tot_el, tot_plain = [], [], []
    for r in sorted(tx):
        for _, t in tx[r]:
            if RE_ASSETS.match(t) and a_hdr is None:
                a_hdr = r
            elif RE_EL.match(t) and l_hdr is None:
                l_hdr = r
            elif RE_TOTAL_A.match(t):
                tot_a.append(r)
            elif RE_TOTAL_EL.match(t):
                tot_el.append(r)
            elif RE_TOTAL.match(t):
                tot_plain.append(r)
    if a_hdr is None or l_hdr is None:
        return None

    def first_after(rows, start, exclude=()):
        for r in sorted(rows):
            if r > start and r not in exclude:
                return r
        return None

    a_tot = first_after(tot_a, a_hdr)
    l_tot = first_after(tot_el, l_hdr)
    if a_tot is None:
        a_tot = first_after(tot_plain, a_hdr, exclude=(l_tot,))
    if l_tot is None:
        l_tot = first_after(tot_plain, l_hdr, exclude=(a_tot,))
    if a_tot is None or l_tot is None:
        return None
    band_end = min(a_hdr, l_hdr)
    ws = bk.wv[sh]
    cols = []
    for c in range(1, min(bk.maxcol(sh), 60) + 1):
        va, vl = ws.cell(a_tot, c).value, ws.cell(l_tot, c).value
        if not ((num(va) is not None or is_err(va)) and (num(vl) is not None or is_err(vl))):
            continue
        heads = []
        for r in range(1, band_end):
            v = ws.cell(r, c).value
            if isinstance(v, (dt.datetime, dt.date)):
                heads.append(v.strftime("%d.%m.%Y"))
            elif isinstance(v, str) and RE_PERIOD.search(v):
                heads.append(clean(v))
        cols.append({"col": c, "letter": get_column_letter(c), "header": " ".join(heads)})
    with_head = [c for c in cols if c["header"]]
    cols = with_head if with_head else cols
    units = ""
    for r in sorted(tx):
        if r > 15:
            break
        for _, t in tx[r]:
            m = re.search(r"(lakh|lac|thousand|crore|million|rupee)", t, re.I)
            if m and re.search(r"\bin\b", t, re.I):
                units = m.group(1).lower().replace("lac", "lakh") + ("s" if not m.group(1).lower().endswith("s") else "")
    return {"sheet": sh, "a_hdr": a_hdr, "a_tot": a_tot, "l_hdr": l_hdr, "l_tot": l_tot, "cols": cols, "units": units}


def candidate_sheets(bk):
    out = []
    for sh in bk.sheetnames:
        try:
            st = find_structure(bk, sh)
        except Exception:
            st = None
        if st and st["cols"]:
            out.append(sh)
    return out


def detect_tb_sheets(bk):
    res = []
    for sh in bk.sheetnames:
        nm = sh.lower()
        hit = bool(re.search(r"(^|[^a-z])tb([^a-z]|$)|trial", nm))
        if not hit:
            ws = bk.wv[sh]
            texts = []
            for row in ws.iter_rows(min_row=1, max_row=min(15, ws.max_row or 1), max_col=min(ws.max_column or 1, 40), values_only=True):
                texts.extend(v.lower() for v in row if isinstance(v, str))
            j = " | ".join(texts)
            hit = "trial balance" in j or ("debit" in j and "credit" in j)
        if hit:
            res.append(sh)
    return res


def label_for_row(bk, sh, r, max_col=6):
    best = ""
    for v in bk.wv[sh].iter_rows(min_row=r, max_row=r, max_col=max_col, values_only=True).__next__():
        if isinstance(v, str):
            t = clean(v)
            if t and not RE_LABEL_JUNK.match(t) and len(t) > len(best):
                best = t
    return best


def classify(bk, sh, col, raw):
    if not is_formula(raw):
        return "hardcode" if num(raw) is not None else "other"
    items = parse_formula(raw) or []
    ops = []
    collect_operands(items, ops)
    rng = [o for o in ops if o.kind == "RANGE"]
    if not rng:
        return "hardcode"
    for o in rng:
        if "!" in o.value:
            return "link"
        for m in re.finditer(r"\$?([A-Z]{1,3})\$?\d+", o.value.upper()):
            if column_index_from_string(m.group(1)) != col:
                return "link"
    return "subtotal"


def get_lines(bk, st, col):
    sh = st["sheet"]
    lines = []
    for side, lo, hi in (("A", st["a_hdr"], st["a_tot"]), ("L", st["l_hdr"], st["l_tot"])):
        for r in range(lo + 1, hi):
            a = addr(col, r)
            raw, cv = bk.raw(sh, a), bk.val(sh, a)
            if raw is None and cv is None:
                continue
            kind = classify(bk, sh, col, raw)
            if kind == "other":
                continue
            lines.append({"row": r, "side": side, "label": label_for_row(bk, sh, r) or f"row {r}", "raw": raw,
                          "value": num(cv), "err": cv if is_err(cv) else None, "kind": kind, "addr": a})
    return lines


# --------------------------------------------------------------------------- tracing
def trace_line(bk, tb_sheets, sh, a):
    seen, depth, scal = {}, {}, {}
    parents = defaultdict(set)
    tb_leaves, hard, errs = {}, [], []
    stack = [((sh, a), 1, None, 0, False)]
    while stack:
        cell, sg, par, d, sc = stack.pop()
        if par is not None:
            parents[cell].add(par)
        if cell in seen:
            depth[cell] = min(depth[cell], d)
            scal[cell] = scal[cell] and sc
            continue
        seen[cell] = sg
        depth[cell] = d
        scal[cell] = sc
        csh, ca = cell
        if csh in tb_sheets and par is not None:
            tb_leaves[cell] = sg
            continue
        raw, cv = bk.raw(csh, ca), bk.val(csh, ca)
        if is_err(cv):
            errs.append(cell)
        if is_formula(raw):
            ch = bk.expand(csh, ca)
            if not ch:
                hard.append((cell, cv))
            scl = bk._scl.get((csh, ca), {})
            for (c2, a2, s2) in ch:
                stack.append(((c2, a2), sg * s2, cell, d + 1, sc or scl.get((c2, a2, s2), False)))
        elif num(cv) not in (None, 0.0):
            hard.append((cell, cv))
    return {"seen": seen, "depth": depth, "parents": parents, "tb": tb_leaves, "hard": hard, "errs": errs, "scal": scal}


def col_role(bk, sh, col, depth=0):
    ws = bk.wv[sh]
    for r in range(1, 16):
        v = ws.cell(r, col).value
        if isinstance(v, str):
            s = v.lower()
            if "debit" in s or re.search(r"\bdr\b", s):
                return "D"
            if "credit" in s or re.search(r"\bcr\b", s):
                return "C"
    if depth < 2:
        wsf = bk.wf[sh]
        for r in range(1, min(bk.maxrow(sh), 80) + 1):
            raw = wsf.cell(r, col).value
            if is_formula(raw):
                for m in re.finditer(r"(?<![!A-Za-z'])\$?([A-Z]{1,3})\$?\d+", raw):
                    c2 = column_index_from_string(m.group(1))
                    if c2 != col:
                        role = col_role(bk, sh, c2, depth + 1)
                        if role:
                            return role
                break
    return None


def aggregates_column(raw, col):
    for m in re.finditer(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)", raw.upper()):
        if column_index_from_string(m.group(1)) == col == column_index_from_string(m.group(3)) and int(m.group(4)) - int(m.group(2)) >= 2:
            return True
    return False


def tb_label(bk, sh, r):
    ws = bk.wv[sh]
    text, code = "", None
    for c, v in enumerate(next(ws.iter_rows(min_row=r, max_row=r, max_col=6, values_only=True)), 1):
        if isinstance(v, str) and clean(v) and not text and not re.match(r"^[\d.]+$", clean(v)):
            text = clean(v)
        elif isinstance(v, (int, float)) and code is None and c == 2 and float(v) == int(v) and 0 < v < 10 ** 7:
            code = int(v)
    return text, code


# --------------------------------------------------------------------------- the main analysis of one gap column
def effect(side, c):
    """effect on (Total E&L - Total Assets) of a contribution c to a line on that side"""
    return c if side == "L" else -c


def analyze_gap_column(bk, st, colinfo, gap, tb_sheets, tol):
    sh, col = st["sheet"], colinfo["col"]
    L = colinfo["letter"]
    lines = get_lines(bk, st, col)
    leaf = [x for x in lines if x["kind"] in ("link", "hardcode")]
    res = {"col": L, "gap": gap, "items": [], "observations": [], "lines": lines}

    # --- internal arithmetic
    for side, tot_row in (("A", st["a_tot"]), ("L", st["l_tot"])):
        tot = num(bk.val(sh, addr(col, tot_row)))
        ls = [x for x in leaf if x["side"] == side]
        if tot is not None and all(x["value"] is not None or x["err"] is None for x in ls):
            s = sum(x["value"] or 0 for x in ls)
            if abs(tot - s) > max(0.005, 1e-6 * abs(tot)):
                e = effect(side, tot - s)
                res["items"].append({"tier": 1, "type": "INTERNAL", "e": e, "side": side,
                                     "text": f"The {'assets' if side == 'A' else 'equity & liabilities'} total ({sh}!{L}{tot_row} = {fmt(tot)}) is not the sum of its lines ({fmt(s)}); difference {fmt(tot - s)}. A line is left out of, or counted twice in, the total formula.",
                                     "ref": f"{sh}!{L}{tot_row}"})

    # --- trace every leaf line to the TB
    traces = {}
    for x in leaf:
        if x["kind"] == "link":
            traces[x["row"]] = trace_line(bk, tb_sheets, sh, x["addr"])
    sheet_count = Counter()
    for t in traces.values():
        for (csh, _) in t["tb"]:
            sheet_count[csh] += 1
    res["tb_reached"] = dict(sheet_count)
    line_by_row = {x["row"]: x for x in lines}

    # --- hard-coded statement lines, negative balances
    for x in leaf:
        if x["kind"] == "hardcode" and x["value"] and abs(x["value"]) > tol:
            res["observations"].append({"sev": "Medium", "type": "HARDCODE",
                                        "text": f"'{x['label']}' ({sh}!{x['addr']}) is typed in as {fmt(x['value'])} instead of being linked."})
        if x["value"] is not None and x["value"] < -tol and x["side"] == "L" and not re.search(r"equity|reserve|surplus|deficit|retained", x["label"], re.I):
            res["observations"].append({"sev": "Info", "type": "NEGATIVE",
                                        "text": f"'{x['label']}' ({sh}!{x['addr']}) shows a negative balance of {fmt(x['value'])} on the liabilities side - a source ledger probably carries the opposite (debit) balance and may need reclassifying."})
        if x["value"] is not None and x["value"] < -tol and x["side"] == "A":
            res["observations"].append({"sev": "Info", "type": "NEGATIVE",
                                        "text": f"'{x['label']}' ({sh}!{x['addr']}) shows a negative balance of {fmt(x['value'])} on the assets side - check for a credit balance that should be shown under liabilities."})

    # --- errors inside the chains
    err_origin = []
    for row, t in traces.items():
        for cell in t["errs"]:
            kids = bk.expand(*cell)
            if not any(is_err(bk.val(k[0], k[1])) for k in kids):
                err_origin.append((cell, line_by_row[row]["label"]))
    for (csh, ca), lab in err_origin[:5]:
        res["observations"].append({"sev": "High", "type": "ERROR",
                                    "text": f"{bk.val(csh, ca)} originates at {csh}!{ca} (formula: {bk.raw(csh, ca)}), feeding '{lab}'."})

    # --- typed-in numbers (statement lines and inputs inside the chains)
    hards = {}
    for row, t in traces.items():
        for cell, v in t["hard"]:
            if cell[0] != sh and num(v) and abs(num(v)) >= max(tol, 0.001 * abs(gap)):
                d0 = t["depth"].get(cell, 99)
                if cell not in hards or d0 < hards[cell][2]:
                    hards[cell] = (num(v), line_by_row[row]["label"], d0)
    typed = {}
    for x in leaf:
        if x["kind"] == "hardcode" and x["value"] and abs(x["value"]) >= 1:
            typed.setdefault(round(abs(x["value"]), 2), f"{sh}!{x['addr']}")
    for cell, (v, lab, _d) in hards.items():
        if abs(v) >= 1:
            typed.setdefault(round(abs(v), 2), f"{cell[0]}!{cell[1]}")

    # --- trial balance coverage
    if sheet_count:
        tb = sheet_count.most_common(1)[0][0]
        res["tb"] = tb
        _tb_coverage(bk, st, colinfo, tb, traces, line_by_row, res, tol, typed)
    else:
        res["tb"] = None
        res["observations"].append({"sev": "Medium", "type": "NOTB",
                                    "text": "No trial-balance sheet was reached from the statement lines, so ledger-level coverage could not be checked. Pick the TB sheet(s) manually if they were not detected."})

    # --- the same line pulled from two period columns of one sheet
    _overlaps(bk, st, colinfo, tb_sheets, traces, line_by_row, res)

    # --- hard-coded inputs inside the chains (information)
    scale = max(abs(num(bk.val(sh, addr(col, st["a_tot"]))) or 0), 1.0)   # drop numbers in a different unit (e.g. rupees vs lakhs)
    res["hardcoded_inputs"] = [{"cell": f"{c[0]}!{c[1]}", "value": v[0], "line": v[1]}
                               for c, v in sorted(hards.items(), key=lambda kv: (kv[1][2], -abs(kv[1][0]))) if abs(v[0]) <= scale][:25]

    res["observations"] += typed_outliers(bk, traces, st, tb_sheets)

    # --- find the combination that explains the gap
    res["solutions"] = solve(gap, res["items"], tol)
    best = res["solutions"][0] if res["solutions"] else None
    if res.get("tb") and (best is None or not best["within"]):
        try:
            tie = tie_out(bk, st, colinfo, tb_sheets, traces, res, gap, tol)
        except Exception as ex:  # keep the rest of the analysis even if the tie-out fails
            tie = None
            bk.warn(f"Ledger tie-out could not be completed: {ex}")
        if tie:
            res["tieout"] = tie
            if tie["items"]:
                res["items"] += tie["items"]
                res["solutions"] = solve(gap, res["items"], tol)
    return res


def _tb_coverage(bk, st, colinfo, tb, traces, line_by_row, res, tol, typed=None):
    typed = typed or {}
    sh = st["sheet"]
    leaves = defaultdict(dict)          # (col,row) -> {line_row: contribution}
    cell_via = defaultdict(set)         # row -> set(parent cells)
    amount_cols = set()
    for row, t in traces.items():
        for (csh, ca), sg in t["tb"].items():
            if csh != tb:
                continue
            c, r = split_addr(ca)
            amount_cols.add(c)
            v = num(bk.val(csh, ca)) or 0.0
            leaves[r][row] = leaves[r].get(row, 0.0) + sg * v
            for p in t["parents"].get((csh, ca), ()):
                cell_via[r].add(f"{p[0]}!{p[1]}")
    amount_cols = sorted(amount_cols)
    roles = {c: col_role(bk, tb, c) for c in amount_cols}
    if None in roles.values():
        for i, c in enumerate(amount_cols):
            if not roles[c]:
                roles[c] = "D" if (i == 0 or "D" not in roles.values()) else "C"
        res["observations"].append({"sev": "Info", "type": "ROLE", "text": f"Debit/credit side of some amount columns in '{tb}' could not be read from headers; assumed {', '.join(get_column_letter(c) + '=' + roles[c] for c in amount_cols)}."})
    # unit scale of each amount column (e.g. a rupee column read as /1000) and mirrored columns (helper = raw/1000)
    scales = defaultdict(list)
    for t in traces.values():
        for (csh, ca), sg in t["tb"].items():
            if csh == tb and not t["scal"].get((csh, ca)):
                scales[split_addr(ca)[0]].append(abs(sg))
    col_scale = {c: (sorted(v)[len(v) // 2] if v else 1.0) for c, v in scales.items()}
    for c in amount_cols:
        col_scale.setdefault(c, 1.0)
    wsf0 = bk.wf[tb]
    parent = {c: c for c in amount_cols}

    def find(c):
        while parent[c] != c:
            c = parent[c]
        return c

    for c in amount_cols:
        for r in list(leaves)[:60]:
            raw = wsf0.cell(r, c).value
            if not is_formula(raw):
                continue
            ops = []
            collect_operands(parse_formula(raw) or [], ops)
            rng = [o for o in ops if o.kind == "RANGE"]
            if len(rng) == 1 and "!" not in rng[0].value:
                m = re.match(r"^\$?([A-Z]{1,3})\$?(\d+)$", rng[0].value.upper())
                if m and int(m.group(2)) == r and column_index_from_string(m.group(1)) in parent and column_index_from_string(m.group(1)) != c:
                    parent[find(c)] = find(column_index_from_string(m.group(1)))
            break
    groups = defaultdict(list)
    for c in amount_cols:
        groups[find(c)].append(c)
    group_rep = {g: max(cs, key=lambda c: (col_scale[c], c)) for g, cs in groups.items()}
    group_role = {g: roles[group_rep[g]] for g in groups}
    res["tb_cols"] = {"+".join(get_column_letter(c) for c in cs): group_role[g] for g, cs in groups.items()}

    wsv, wsf = bk.wv[tb], bk.wf[tb]
    rows = {}
    for r in range(1, bk.maxrow(tb) + 1):
        raw_vals = {c: num(wsv.cell(r, c).value) for c in amount_cols}
        if all(v is None for v in raw_vals.values()):
            continue
        gvals = {}
        for g, cs in groups.items():
            rep = group_rep[g]
            v = raw_vals.get(rep)
            if v is None:
                alt = [(raw_vals[c], col_scale[c]) for c in cs if raw_vals.get(c) is not None]
                gvals[g] = alt[0][0] * alt[0][1] if alt else None
            else:
                gvals[g] = v * col_scale[rep]
        label, code = tb_label(bk, tb, r)
        sub = bool(re.match(r"^\W*(grand\s+)?total", label, re.I))
        for c in amount_cols:
            raw = wsf.cell(r, c).value
            if is_formula(raw) and aggregates_column(raw, c):
                sub = True
        dr = sum(v or 0 for g, v in gvals.items() if group_role[g] == "D")
        cr = sum(v or 0 for g, v in gvals.items() if group_role[g] == "C")
        rows[r] = {"row": r, "label": label, "code": code, "gvals": gvals, "dr": dr, "cr": cr, "net": dr - cr, "sub": sub}
    res["tb_rows_total"] = sum(1 for x in rows.values() if not x["sub"])

    linked_rows = {r for r in leaves}
    # neighbours (for "probably belongs to ..." hints)
    def hint(r):
        near = [q for q in range(r - 3, r + 4) if q != r and q in linked_rows and q in rows and not rows[q]["sub"]]
        labs = Counter(line_by_row[k]["label"] for q in near for k in leaves[q])
        if not near or not labs:
            return ""
        lab, n = labs.most_common(1)[0]
        if n >= max(2, len(near) // 2):
            return lab
        return ""

    tb_unlinked, tb_multi, tb_partial = [], [], []
    for r, x in rows.items():
        if x["sub"]:
            continue
        mag = max([abs(v) for v in x["gvals"].values() if v is not None] or [0.0])
        if r not in linked_rows:
            if mag > 1e-9:
                h = hint(r)
                ty = typed.get(round(abs(x["net"]), 2), "")
                tb_unlinked.append({"row": r, "label": x["label"], "code": x["code"], "net": x["net"], "dr": x["dr"], "cr": x["cr"], "hint": h, "typed_in": ty})
                res["items"].append({"tier": 2, "type": "UNLINKED", "e": x["net"], "ref": f"{tb} row {r}",
                                     "text": f"TB '{tb}' row {r} - {x['label'] or 'unnamed ledger'}{'' if x['code'] is None else ' (code ' + str(x['code']) + ')'}: {('Dr ' + fmt(x['dr'])) if x['dr'] else ('Cr ' + fmt(x['cr']))} is not read by any balance sheet line" + (f"; its neighbouring ledgers feed '{h}', so it probably belongs there" if h else "") + (f". Its amount equals the typed-in number at {ty} - the figure was probably typed in instead of linked" if ty else "") + "."})
        else:
            # partly linked: an amount column with a value that no line reads
            reached_cols = {split_addr(ca)[0] for row_ in traces.values() for (csh, ca) in row_["tb"] if csh == tb and split_addr(ca)[1] == r}
            reached_groups = {find(c) for c in reached_cols if c in parent}
            miss = [g for g in groups if g not in reached_groups and abs(x["gvals"].get(g) or 0) > 1e-9]
            if miss and reached_groups:
                net = sum((x["gvals"][g] or 0) * (1 if group_role[g] == "D" else -1) for g in miss)
                names = ["+".join(get_column_letter(c) for c in groups[g]) for g in miss]
                tb_partial.append({"row": r, "label": x["label"], "cols": names, "net": net})
                res["items"].append({"tier": 2, "type": "PARTIAL", "e": net, "ref": f"{tb} row {r}",
                                     "text": f"TB '{tb}' row {r} - {x['label']}: column(s) {', '.join(names)} hold {fmt(abs(net))} that no line reads, although the other side of this ledger is linked (e.g. a SUMIF without the matching debit/credit deduction, or a link that takes only debit or only credit)."})
    for d in _dup_events(bk, st, colinfo, tb, traces, line_by_row):
        tb_multi.append(d)
        res["items"].append({"tier": 1, "type": "DUP", "e": d["e"], "ref": f"{tb} row {d['row']}", "text": d["text"]})
    res["tb_unlinked"] = sorted(tb_unlinked, key=lambda d: -abs(d["net"]))
    res["tb_multi"] = tb_multi
    res["tb_partial"] = tb_partial
    res["tb_rows_linked"] = len(linked_rows)


def _dup_events(bk, st, colinfo, tb, traces, line_by_row):
    """A ledger amount that is linked from two (or more) different cells/sub-heads on the same side, with the same sign,
    is counted more than once. Works for SUMIF and direct links, inside one statement line or across lines."""
    events = defaultdict(list)        # tb row -> read events
    for lrow, t in traces.items():
        side = line_by_row[lrow]["side"]
        by_parent = defaultdict(lambda: defaultdict(list))   # row -> parent -> [(cell, c, m)]
        for (csh, ca), _sg in t["tb"].items():
            if csh != tb:
                continue
            c, r = split_addr(ca)
            v = num(bk.val(csh, ca)) or 0.0
            for p in t["parents"].get((csh, ca), ()):
                psign = t["seen"].get(p, 1)
                if t["scal"].get(p):
                    continue                    # parent already reached through a rate / multiplication
                mu = bk._mult.get(p) or Counter()
                sc_map = bk._scl.get(p) or {}
                for (k_sh, k_a, s2), m in mu.items():
                    if (k_sh, k_a) == (csh, ca) and not sc_map.get((k_sh, k_a, s2)):
                        by_parent[r][p].append((ca, psign * s2 * v, m))
        for r, parents in by_parent.items():
            for p, lst in parents.items():
                reads = max(m for _, _, m in lst)
                c_total = sum(cv * m for _, cv, m in lst)
                cells = [x[0] for x in lst]
                for _ in range(reads):
                    events[r].append({"line": lrow, "side": side, "parent": p, "c": c_total / reads, "cells": cells})
    out = []
    for r, evs in sorted(events.items()):
        label, code = tb_label(bk, tb, r)
        for side in ("A", "L"):
            grp = [e for e in evs if e["side"] == side and abs(e["c"]) > 1e-9]
            if len(grp) < 2:
                continue
            total = sum(g["c"] for g in grp)
            unit = max(abs(g["c"]) for g in grp)
            k = int(round(abs(total) / unit)) if unit else 0
            # cancelling pair = opposite-sign reads that land in the SAME statement line (they wipe each other out there);
            # opposite reads in different lines are an ordinary two-sided journal (e.g. expense and its accrued payable)
            cancel = k == 0 and any(
                any(g["c"] > 1e-9 and g["line"] == ln for g in grp) and any(g["c"] < -1e-9 and g["line"] == ln for g in grp)
                for ln in {g["line"] for g in grp})
            if k < 2 and not cancel:
                continue   # extra reads are cancelled by deductions (e.g. non-current = total - current): net count is one
            grp.sort(key=lambda e: (e["line"], f"{e['parent'][0]}!{e['parent'][1]}"))
            where = "; ".join(f"via {g['parent'][0]}!{g['parent'][1]} (feeds '{line_by_row[g['line']]['label']}', {fmt(g['c'])})" for g in grp[:6])
            head = f"TB '{tb}' row {r} - {label or 'ledger'}{'' if code is None else ' (code ' + str(code) + ')'}: the same ledger amount is linked from {len(grp)} different places on the {'assets' if side == 'A' else 'equity & liabilities'} side - {where}. "
            if not cancel:
                e_tot = effect(side, total - math.copysign(unit, total))
                text = head + f"Net, it is counted {k} times ({fmt(total)} instead of {fmt(math.copysign(unit, total))}); link it from one sub-head only, or deduct the portion shown elsewhere."
            else:
                # the link that sits in a formula of its own is taken as the proper one; the link buried in a bigger formula is the stray
                reads = {g["parent"]: len(bk.expand(*g["parent"])) for g in grp}
                proper = min(grp, key=lambda g: (reads[g["parent"]], abs(g["c"]) * -1))
                stray = [g for g in grp if g is not proper] or grp[1:]
                e_tot = effect(side, total - proper["c"])
                text = head + (f"The reads have opposite signs, so the ledger cancels itself out (net {fmt(total)} instead of {fmt(proper['c'])}). "
                               f"The link at {proper['parent'][0]}!{proper['parent'][1]} looks like the intended one; "
                               f"{', '.join(g['parent'][0] + '!' + g['parent'][1] for g in stray[:3])} reads the same cell again and is probably the stray link.")
                for g in stray[:1]:
                    pcell = g["parent"]
                    mine = {(c2, s2) for (c2, a2, s2) in bk.expand(*pcell) if c2 == tb and a2 in g["cells"]}
                    my_sign = next(iter([s2 for (c2, a2, s2) in bk.expand(*pcell) if c2 == tb and a2 in g["cells"]]), None)
                    cols = {split_addr(a2)[0] for a2 in g["cells"]}
                    others = [(a2, s2) for (c2, a2, s2) in bk.expand(*pcell)
                              if c2 == tb and a2 not in g["cells"] and split_addr(a2)[0] in cols and my_sign and s2 * my_sign < 0]
                    if others:
                        text += (f" In that formula the other cells of the same column ({', '.join(a2 for a2, _ in others[:5])}) are "
                                 f"{'deducted' if others[0][1] < 0 else 'added'}, but {', '.join(g['cells'])} is "
                                 f"{'added' if (my_sign or 1) > 0 else 'deducted'} - check the sign.")
            out.append({"row": r, "label": label, "side": side, "lines": [line_by_row[g["line"]]["label"] for g in grp],
                        "contrib": [g["c"] for g in grp], "e": e_tot, "via": [f"{g['parent'][0]}!{g['parent'][1]}" for g in grp], "text": text})
    return out


def scan_tallying_column(bk, st, colinfo, tb_sheets, tol):
    """Light scan for columns that tally: duplicate links can be hidden by an offsetting error."""
    sh = st["sheet"]
    lines = get_lines(bk, st, colinfo["col"])
    line_by_row = {x["row"]: x for x in lines}
    traces = {x["row"]: trace_line(bk, tb_sheets, sh, x["addr"]) for x in lines if x["kind"] == "link"}
    cnt = Counter(csh for t in traces.values() for (csh, _) in t["tb"])
    res = {"col": colinfo["letter"], "items": [], "observations": typed_outliers(bk, traces, st, tb_sheets), "solutions": [], "dup_only": True,
           "tb": None, "tb_multi": []}
    if not cnt:
        return res
    tb = cnt.most_common(1)[0][0]
    res["tb"] = tb
    for d in _dup_events(bk, st, colinfo, tb, traces, line_by_row):
        res["tb_multi"].append(d)
        res["items"].append({"tier": 1, "type": "DUP", "e": d["e"], "ref": f"{tb} row {d['row']}", "text": d["text"]})
    return res


def _overlaps(bk, st, colinfo, tb_sheets, traces, line_by_row, res):
    sh = st["sheet"]
    first_line, sign_of, depth, scaled_cell = {}, {}, {}, {}
    children = {}
    for row, t in traces.items():
        for cell, sg in t["seen"].items():
            if cell not in first_line:
                first_line[cell], sign_of[cell] = row, sg
            scaled_cell[cell] = scaled_cell.get(cell, True) and t["scal"].get(cell, False)
            depth[cell] = min(depth.get(cell, 10 ** 6), t["depth"][cell])
    cells = set(first_line)

    def kids(cell):
        if cell not in children:
            children[cell] = [(a, b) for a, b, _ in bk.expand(*cell)] if cell[0] not in tb_sheets else []
        return children[cell]

    memo = {}

    def desc(cell):
        if cell in memo:
            return memo[cell]
        out, stack = set(), [cell]
        while stack:
            cur = stack.pop()
            for k in kids(cur):
                if k not in out:
                    out.add(k)
                    stack.append(k)
        memo[cell] = out
        return out

    groups = defaultdict(list)
    for cell in cells:
        csh, ca = cell
        if csh in tb_sheets or csh == sh:
            continue
        v = num(bk.val(csh, ca))
        if v is None:
            continue
        c, r = split_addr(ca)
        groups[(csh, r)].append((c, cell, v))
    pairs = []
    for (csh, r), mem in groups.items():
        mem.sort()
        for (c1, cell1, v1), (c2, cell2, v2) in itertools.combinations(mem, 2):
            if abs(v1) < 1e-9 and abs(v2) < 1e-9:
                continue
            if scaled_cell.get(cell1) or scaled_cell.get(cell2):
                continue
            if cell2 in desc(cell1) or cell1 in desc(cell2):
                continue  # roll-forward (opening = previous closing)
            pairs.append((max(depth[cell1], depth[cell2]), cell1, cell2, v1, v2))
    pairs.sort(key=lambda p: p[0])
    covered = set()
    out_pairs = []
    for d, c1, c2, v1, v2 in pairs:
        if c1 in covered and c2 in covered:
            continue
        covered |= desc(c1) | desc(c2) | {c1, c2}
        out_pairs.append((c1, c2, v1, v2))
    reported = 0
    for c1, c2, v1, v2 in out_pairs:
        if reported >= 12:
            break
        ws_sh = c1[0]
        col1, r = split_addr(c1[1])
        col2, _ = split_addr(c2[1])
        hdr = lambda cc: clean(" ".join(str(bk.val(ws_sh, addr(cc, rr)) or "") for rr in range(1, 12) if isinstance(bk.val(ws_sh, addr(cc, rr)), str) and RE_PERIOD.search(str(bk.val(ws_sh, addr(cc, rr))))))
        h1, h2 = hdr(col1), hdr(col2)
        d1, d2 = parse_date(h1), parse_date(h2)
        par1 = sorted({f"{p[0]}!{p[1]}" for t in traces.values() for p in t["parents"].get(c1, ())})[:2]
        par2 = sorted({f"{p[0]}!{p[1]}" for t in traces.values() for p in t["parents"].get(c2, ())})[:2]
        line1, line2 = line_by_row[first_line[c1]], line_by_row[first_line[c2]]
        if not (d1 and d2 and d1 != d2):
            continue  # not period columns (e.g. fixed-asset schedule columns) - not an overlap
        cand = [(c1, v1)] if d1 < d2 else [(c2, v2)]
        txt = (f"{ws_sh}!{c1[1]} ({h1 or 'column ' + get_column_letter(col1)}, {fmt(v1)}) and {ws_sh}!{c2[1]} ({h2 or 'column ' + get_column_letter(col2)}, {fmt(v2)}) are the same row of two different period columns, "
               f"and both feed '{line1['label']}' (through {', '.join(par1)} and {', '.join(par2)}). If the later column is cumulative, the earlier period's amount ({fmt(cand[0][1])}) is counted twice.")
        entry = {"cells": [f"{c1[0]}!{c1[1]}", f"{c2[0]}!{c2[1]}"], "values": [v1, v2], "text": txt}
        reported += 1
        res.setdefault("overlaps", []).append(entry)
        for cell, v in cand:
            ln = line_by_row[first_line[cell]]
            e = effect(ln["side"], sign_of[cell] * v)
            res["items"].append({"tier": 1, "type": "OVERLAP", "e": e, "ref": f"{cell[0]}!{cell[1]}", "text": txt, "short": f"{cell[0]}!{cell[1]} counted twice"})


# --------------------------------------------------------------------------- explaining the gap
def solve(gap, items, tol):
    tol_ok = max(tol, 5e-4 * abs(gap))
    t1 = sorted([i for i in items if i["tier"] == 1 and abs(i["e"]) > 1e-9], key=lambda i: -abs(i["e"]))[:14]
    t2 = [i for i in items if i["tier"] == 2 and abs(i["e"]) > 1e-9 and abs(i["e"]) <= 5 * abs(gap) + tol]
    t2 = sorted(t2, key=lambda i: -abs(i["e"]))[:150]
    n = len(t1)
    subs = sorted(((sum(t1[i]["e"] for i in range(n) if m >> i & 1), m) for m in range(1 << n)))
    sums = [s for s, _ in subs]

    def near(need, k=2):
        pos = bisect.bisect_left(sums, need)
        return [subs[p] for p in range(max(0, pos - k), min(len(subs), pos + k))]

    results = []

    def push(sel2, s2):
        for s, m in near(gap - s2):
            its = [t1[i] for i in range(n) if m >> i & 1] + list(sel2)
            results.append((abs(gap - s - s2), len(its), its))

    push((), 0.0)
    ok1 = [r for r in results if r[0] <= tol_ok and r[1] > 0]
    if not ok1 and t2:
        for k in (1, 2, 3):
            for combo in itertools.combinations(t2, k):
                push(combo, sum(i["e"] for i in combo))
            results.sort(key=lambda r: (r[0] > tol, r[0] > tol_ok, r[1], r[0]))
            del results[200:]
    results = [r for r in results if r[1] > 0]
    results.sort(key=lambda r: (r[0] > tol, r[0] > tol_ok, r[1], r[0]))
    out, seen = [], set()
    for resid, k, its in results:
        key = frozenset(id(i) for i in its)
        if key in seen:
            continue
        seen.add(key)
        out.append({"residual": gap - sum(i["e"] for i in its), "abs_residual": resid, "n": k, "items": its, "exact": resid <= tol,
                    "within": resid <= tol_ok})
        if len(out) >= 5:
            break
    return out



# --------------------------------------------------------------------------- tie-out of unlinked ledgers to typed-in schedules
STOP_WORDS = {"total", "cost", "amount", "amounts", "march", "april", "balance", "description", "note", "year", "charge", "from",
              "with", "that", "this", "other", "account", "accounts", "ledger", "carrying", "value", "rate", "date", "sheet"}


def _stems(text):
    return {w[:5] for w in re.findall(r"[a-z]{4,}", str(text).lower()) if w not in STOP_WORDS}


def cell_context(bk, sh, a):
    """(primary, secondary) words describing a cell: primary = row label + column header, secondary = headings above + sheet name."""
    key = (sh, a)
    if key in bk._ctx:
        return bk._ctx[key]
    ws = bk.wv[sh]
    c, r = split_addr(a)
    prim, sec = set(), set()
    for v in next(ws.iter_rows(min_row=r, max_row=r, max_col=6, values_only=True)):
        if isinstance(v, str):
            prim |= _stems(v)
    for rr in range(r - 1, max(0, r - 46), -1):
        v = ws.cell(rr, c).value
        if isinstance(v, str) and v.strip():
            prim |= _stems(v)
            break
    found = 0
    for rr in range(r - 1, max(0, r - 60), -1):
        row = next(ws.iter_rows(min_row=rr, max_row=rr, max_col=12, values_only=True))
        vals = [v for v in row if v is not None and not (isinstance(v, str) and not v.strip())]
        if len(vals) == 1 and isinstance(vals[0], str):
            sec |= _stems(vals[0])
            found += 1
            if found == 2:
                break
    sec |= _stems(sh)
    bk._ctx[key] = (prim, sec)
    return bk._ctx[key]


def _name_score(bk, names, m):
    """2 points per word shared with a cell's row label / column header, 1 per word shared with its headings"""
    score = 0
    for c in m:
        prim, sec = cell_context(bk, c[0][0], c[0][1])
        score += 2 * len(names & prim) + len(names & sec)
    return score


def _cell_index(bk, traces, st, tb_sheets):
    cells = {}
    for t in traces.values():
        for cell in t["seen"]:
            if cell[0] in tb_sheets or cell[0] == st["sheet"]:
                continue
            v = num(bk.val(*cell))
            if v and abs(v) >= 0.01:
                cells[cell] = v
    vals, metas = [], []
    for cell, v in cells.items():
        for div, tag in ((1.0, ""), (1e3, "/1000"), (1e5, "/10^5")):
            w = abs(v) / div
            if w >= 0.01:
                vals.append(w)
                metas.append((cell, tag, v))
    if not vals:
        return np.array([]), []
    order = np.argsort(vals)
    return np.array(vals)[order], [metas[i] for i in order]


def _match(vals, metas, T, tol, max_out=40):
    """schedule figures (one cell, or the sum of two cells) equal to T"""
    out = []
    lo, hi = np.searchsorted(vals, T - tol, "left"), np.searchsorted(vals, T + tol, "right")
    for i in range(lo, hi):
        out.append((metas[i],))
    need = T - vals
    lo, hi = np.searchsorted(vals, need - tol, "left"), np.searchsorted(vals, need + tol, "right")
    for i in np.nonzero(hi > lo)[0]:
        for j in range(lo[i], hi[i]):
            if j > i and metas[i][0] != metas[j][0] and metas[i][0][0] == metas[j][0][0]:
                ci, ri = split_addr(metas[i][0][1])
                cj, rj = split_addr(metas[j][0][1])
                if ci == cj or ri == rj:            # a real schedule adds figures from the same column or row of one sheet
                    out.append((metas[i], metas[j]))
                    if len(out) >= max_out:
                        return out
    return out


def _cells_text(m):
    return " + ".join(f"{c[0][0]}!{c[0][1]}{(' ' + c[1]) if c[1] else ''}" for c in m)


def _typed_leaves(bk, cell, limit=8):
    out, seen, stack = [], {cell}, [cell]
    while stack and len(seen) < 3000:
        cur = stack.pop()
        raw = bk.raw(*cur)
        if is_formula(raw):
            kids = bk.expand(*cur)
            if not kids:
                v = num(bk.val(*cur))
                if v:
                    out.append((cur, v))
            for (a, b, _s) in kids:
                if (a, b) not in seen:
                    seen.add((a, b))
                    stack.append((a, b))
        elif num(raw):
            out.append((cur, num(raw)))
    out.sort(key=lambda x: -abs(x[1]))
    return [{"cell": f"{c[0]}!{c[1]}", "value": v, "label": label_for_row(bk, c[0], split_addr(c[1])[1]) or ""} for c, v in out[:limit]]


def tie_out(bk, st, colinfo, tb_sheets, traces, res, gap, tol):
    """Ledgers that no statement line reads are compared, by name and amount, with the typed-in / schedule figures."""
    rows = [u for u in res.get("tb_unlinked", []) if abs(u["net"]) > 0.0005]
    if not rows:
        return None
    vals, metas = _cell_index(bk, traces, st, tb_sheets)
    if len(vals) == 0:
        return None
    rows = sorted(rows, key=lambda u: -abs(u["net"]))[:16]
    n = len(rows)
    stems_row = [_stems(u["label"]) for u in rows]
    TOL = 0.0007  # trial-balance amounts carry three decimals when shown in thousands
    tied = []
    for k in (1, 2, 3):
        for sub in itertools.combinations(range(n), k):
            T = abs(sum(rows[i]["net"] for i in sub))
            if T < max(1.0, 0.05 * sum(abs(rows[i]["net"]) for i in sub)):
                continue        # tiny totals / heavy netting of opposite balances match random figures
            names = set().union(*[stems_row[i] for i in sub])
            ms = _match(vals, metas, T, TOL)
            if not ms:
                continue
            best = max(((_name_score(bk, names, m), -len(m), m) for m in ms), key=lambda x: (x[0], x[1]))
            sc, _, m = best
            # accepted when a name backs it up (row label / column header), or when several ledgers add up exactly to one typed figure
            if sc >= 2 or (len(sub) >= 2 and len(m) == 1):
                tied.append((sc >= 2, len(sub), sub, m))
    tied.sort(key=lambda x: (not x[0], -x[1]))
    used, verified = set(), []
    for named, sz, sub, m in tied:
        if used & set(sub):
            continue
        used |= set(sub)
        verified.append({"rows": [rows[i]["label"] for i in sub], "total": abs(sum(rows[i]["net"] for i in sub)),
                         "cells": _cells_text(m), "name_confirmed": named})
    unverified = [i for i in range(n) if i not in used]
    cands = []
    if abs(gap) > tol and unverified:
        for k in (1, 2, 3):
            for sub in itertools.combinations(unverified, k):
                s_net = sum(rows[i]["net"] for i in sub)
                T = abs(s_net)
                if T < max(1.0, 0.05 * sum(abs(rows[i]["net"]) for i in sub)):
                    continue
                S = T - gap if s_net > 0 else T + gap
                if S <= 0.01:
                    continue
                names = set().union(*[stems_row[i] for i in sub])
                for m in _match(vals, metas, S, max(TOL, tol / 2), max_out=12):
                    sc = _name_score(bk, names, m)
                    if sc >= 2:
                        cands.append((-sc, len(m), len(sub), sub, m, S, T, s_net))
    cands.sort(key=lambda x: x[:3])
    items, shown, by_key = [], [], {}
    for _, _, _, sub, m, S, T, s_net in cands:
        key = (sub, round(S, 3))
        if key in by_key:
            by_key[key]["also"].append(_cells_text(m))
            continue
        if len(shown) >= 2:
            continue
        D = S - T
        e = -D if s_net > 0 else D
        labs = ", ".join(f"{rows[i]['label']} ({fmt(abs(rows[i]['net']))})" for i in sub)
        text = (f"Ledger(s) not read by any line - {labs} - total {fmt(T)} in the trial balance, but the schedule figure {_cells_text(m)} shows {fmt(S)}: "
                f"{fmt(abs(D))} {'more' if D > 0 else 'less'} than the ledger. That difference equals the gap, so the typed-in inputs or formulas behind "
                f"{_cells_text(m)} are the likely source (e.g. an addition, disposal or depreciation figure entered wrongly).")
        entry = {"rows": [rows[i]["label"] for i in sub], "ledger_total": T, "schedule_total": S, "cells": _cells_text(m), "difference": D,
                 "inputs": _typed_leaves(bk, m[0][0]), "also": []}
        by_key[key] = entry
        shown.append(entry)
        items.append({"tier": 1, "type": "TIEOUT", "e": e, "ref": _cells_text(m), "text": text, "short": f"schedule {_cells_text(m)} vs ledger"})
    return {"verified": verified, "unverified": [{"label": rows[i]["label"], "net": rows[i]["net"]} for i in unverified],
            "candidates": shown, "items": items}



def typed_outliers(bk, traces, st, tb_sheets):
    """A typed number that differs from the rest of its column block (e.g. one depreciation rate 0.73 among ten 0.63)."""
    cols = defaultdict(list)
    for t in traces.values():
        for cell in t["seen"]:
            if cell[0] in tb_sheets or cell[0] == st["sheet"]:
                continue
            raw = bk.raw(*cell)
            if is_formula(raw) or num(raw) in (None, 0.0):
                continue
            c, r = split_addr(cell[1])
            cols[(cell[0], c)].append((r, float(raw)))
    out = []
    for (sh, c), lst in cols.items():
        lst = sorted(set(lst))
        runs, cur = [], [lst[0]] if lst else []
        for item in lst[1:]:
            if item[0] - cur[-1][0] <= 3:
                cur.append(item)
            else:
                runs.append(cur)
                cur = [item]
        if cur:
            runs.append(cur)
        for run in runs:
            if len(run) < 6:
                continue
            cnt = Counter(round(v, 6) for _, v in run)
            mode, k = cnt.most_common(1)[0]
            odd = [(r, v) for r, v in run if round(v, 6) != mode]
            if k >= 0.7 * len(run) and 1 <= len(odd) <= 2:
                for r, v in odd:
                    out.append({"sev": "Medium", "type": "OUTLIER",
                                "text": f"{sh}!{get_column_letter(c)}{r} is typed as {fmt(v, 4)} while {k} of the other {len(run) - len(odd)} typed cells in "
                                        f"{sh} column {get_column_letter(c)} (rows {run[0][0]}-{run[-1][0]}) hold {fmt(mode, 4)} - check whether this input is a slip "
                                        f"(for example a rate or an amount typed differently from its peers)."})
    return out

# --------------------------------------------------------------------------- formulas that differ between period columns
def _shape(raw, col):
    parts = []
    for t in Tokenizer(raw).items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" not in t.value:
            parts.append(re.sub(r"\$?([A-Z]{1,3})\$?(\d+)",
                                lambda m: f"C[{column_index_from_string(m.group(1)) - col}]R{m.group(2)}", t.value.upper()))
        else:
            parts.append(t.value)
    return "".join(parts)


def _rows_in(raw):
    rows = set()
    for m in re.finditer(r"\$?[A-Z]{1,3}\$?(\d+)(?::\$?[A-Z]{1,3}\$?(\d+))?", raw.upper()):
        a = int(m.group(1))
        b = int(m.group(2)) if m.group(2) else a
        rows.update(range(min(a, b), max(a, b) + 1))
    return rows


def shape_checks(bk, st, tol):
    sh = st["sheet"]
    cols = st["cols"]
    out = []
    if len(cols) < 3:
        return out
    rows = set(range(st["a_hdr"] + 1, st["a_tot"] + 1)) | set(range(st["l_hdr"] + 1, st["l_tot"] + 1))
    for r in sorted(rows):
        shapes = {}
        for ci in cols:
            raw = bk.raw(sh, addr(ci["col"], r))
            if is_formula(raw) and classify(bk, sh, ci["col"], raw) == "subtotal":
                try:
                    shapes[ci["col"]] = (_shape(raw, ci["col"]), raw)
                except Exception:
                    pass
        if len(shapes) < 3:
            continue
        cnt = Counter(s for s, _ in shapes.values())
        if len(cnt) < 2:
            continue
        (dom, n_dom), = cnt.most_common(1)
        if list(cnt.values()).count(n_dom) > 1:
            continue
        dom_col = next(c for c, (s, _) in shapes.items() if s == dom)
        dom_rows = _rows_in(shapes[dom_col][1])
        for c, (s, raw) in shapes.items():
            if s == dom:
                continue
            mine = _rows_in(raw)
            # compare with the dominant formula shifted to this column
            dom_raw = shapes[dom_col][1]
            diff = (mine ^ dom_rows)
            impact = 0.0
            for rr in diff:
                v = num(bk.val(sh, addr(c, rr)))
                if v and rr != r:
                    impact += abs(v)
            L = get_column_letter(c)
            miss = sorted(dom_rows - mine)
            extra = sorted(mine - dom_rows)
            what = []
            if miss:
                what.append("leaves out row(s) " + ", ".join(map(str, miss)))
            if extra:
                what.append("adds row(s) " + ", ".join(map(str, extra)))
            out.append({"sev": "High" if impact > tol else "Low", "col": L, "row": r,
                        "text": f"{sh}!{L}{r} ({raw}) is built differently from the same row in {n_dom} other period column(s) (e.g. {raw_example(shapes, dom)}) - it {' and '.join(what) or 'uses a different layout'}"
                                + (f"; the rows involved carry {fmt(impact)} in this column, so this changes the total." if impact > tol else "; the rows involved are zero in this column, so there is no effect today, but it should be made consistent.")})
    return out


def raw_example(shapes, dom):
    for c, (s, raw) in shapes.items():
        if s == dom:
            return raw
    return ""


# --------------------------------------------------------------------------- top level
def pick_default_sheet(bk, cands):
    """Prefer the statement whose newest period is the latest (then names that start with BS)."""
    best, best_key = None, None
    for sh in cands:
        st = find_structure(bk, sh)
        dates = [parse_date(c["header"]) for c in st["cols"] if c["header"]]
        dates = [d for d in dates if d]
        key = (max(dates) if dates else dt.date.min, bool(re.match(r"^\s*bs\b", sh, re.I)), -len(sh))
        if best_key is None or key > best_key:
            best, best_key = sh, key
    return best


def load_book(data: bytes, name="workbook.xlsx"):
    return Book(data, name)


def analyze(bk: Book, sheet=None, tb_sheets=None, tol=0.01, scan_all=True):
    cands = candidate_sheets(bk)
    if sheet is None:
        sheet = pick_default_sheet(bk, cands)
    if not sheet:
        return {"ok": False, "error": "Could not find a balance sheet (a sheet with 'ASSETS', 'EQUITY AND LIABILITIES' and total rows).",
                "candidates": cands}
    st = find_structure(bk, sheet)
    if not st:
        return {"ok": False, "error": f"Sheet '{sheet}' does not look like a balance sheet.", "candidates": cands}
    tbs = list(tb_sheets) if tb_sheets else detect_tb_sheets(bk)
    tbs = [s for s in tbs if s != sheet]
    result = {"ok": True, "file": bk.name, "sheet": sheet, "units": st["units"], "tol": tol, "tb_sheets": tbs,
              "candidates": cands, "structure": {k: st[k] for k in ("a_hdr", "a_tot", "l_hdr", "l_tot")},
              "columns": [], "analyses": {}, "observations": []}
    for ci in st["cols"]:
        ta, tl = bk.val(sheet, addr(ci["col"], st["a_tot"])), bk.val(sheet, addr(ci["col"], st["l_tot"]))
        row = {"col": ci["letter"], "period": ci["header"] or f"column {ci['letter']}", "assets": num(ta), "el": num(tl),
               "ta_cell": f"{sheet}!{ci['letter']}{st['a_tot']}", "tl_cell": f"{sheet}!{ci['letter']}{st['l_tot']}"}
        if is_err(ta) or is_err(tl):
            row["status"], row["gap"] = "Error", None
            row["error_cells"] = [c for c, v in ((row["ta_cell"], ta), (row["tl_cell"], tl)) if is_err(v)]
        else:
            row["gap"] = (tl or 0) - (ta or 0)
            row["status"] = "Tallies" if abs(row["gap"]) <= tol else "Gap"
        result["columns"].append(row)
    for ci, row in zip(st["cols"], result["columns"]):
        if row["status"] == "Gap":
            result["analyses"][ci["letter"]] = analyze_gap_column(bk, st, ci, row["gap"], set(tbs), tol)
        elif row["status"] == "Error":
            result["analyses"][ci["letter"]] = _error_column(bk, st, ci, set(tbs), tol)
        elif row["status"] == "Tallies" and scan_all:
            r_ = scan_tallying_column(bk, st, ci, set(tbs), tol)
            if r_["items"] or r_["observations"]:
                result["analyses"][ci["letter"]] = r_
    result["observations"] = shape_checks(bk, st, tol)
    result["warnings"] = list(bk.warnings)
    result["markdown"] = build_markdown(result)
    return result


def _error_column(bk, st, ci, tbs, tol):
    sh = st["sheet"]
    lines = get_lines(bk, st, ci["col"])
    res = {"col": ci["letter"], "items": [], "observations": [], "lines": lines, "solutions": [], "error_column": True}
    origins = {}
    for x in lines:
        if x["kind"] == "link" and (x["err"] or True):
            t = trace_line(bk, tbs, sh, x["addr"])
            for cell in t["errs"]:
                kids = bk.expand(*cell)
                raw = bk.raw(*cell)
                if not any(is_err(bk.val(k[0], k[1])) for k in kids) or (is_formula(raw) and "#REF!" in raw):
                    origins.setdefault(cell, x["label"])
    for x in lines:
        if x["err"]:
            res["observations"].append({"sev": "High", "type": "ERROR", "text": f"{sh}!{x['addr']} '{x['label']}' shows {x['err']}."})
    for (csh, ca), lab in list(origins.items())[:6]:
        res["observations"].append({"sev": "High", "type": "ERROR", "text": f"The error starts at {csh}!{ca} (formula: {bk.raw(csh, ca)}), which feeds '{lab}'."})
    for side, tot_row in (("A", st["a_tot"]), ("L", st["l_tot"])):
        raw = bk.raw(sh, addr(ci["col"], tot_row))
        if is_formula(raw):
            lit = {x["row"] for x in lines if x["side"] == side and x["kind"] in ("link", "hardcode")}
            have = _rows_in(raw)
            if not lit <= have and not re.search(r"\bSUM\b", raw, re.I):
                pass
    return res


# --------------------------------------------------------------------------- the written analysis
def _side_word(gap):
    return "equity & liabilities are higher than assets" if gap > 0 else "assets are higher than equity & liabilities"


def build_markdown(res):
    u = f" (amounts in {res['units']})" if res["units"] else ""
    md = [f"# Balance sheet gap analysis{u}", f"**Workbook:** {res['file']}  |  **Balance sheet sheet:** {res['sheet']}  |  **Trial balance sheets detected:** {', '.join(res['tb_sheets']) or 'none found'}", ""]
    cols = res["columns"]
    bad = [c for c in cols if c["status"] == "Gap"]
    err = [c for c in cols if c["status"] == "Error"]
    ok = [c for c in cols if c["status"] == "Tallies"]
    parts = []
    if ok:
        parts.append(f"{len(ok)} period column(s) tally ({', '.join(c['col'] for c in ok)})")
    if bad:
        parts.append(f"{len(bad)} do not tally ({', '.join(c['col'] for c in bad)})")
    if err:
        parts.append(f"{len(err)} contain formula errors ({', '.join(c['col'] for c in err)})")
    md.append("**Result:** " + "; ".join(parts) + ".")
    md += ["", "| Column | Period | Total Assets | Total Equity & Liabilities | Difference | Status |", "|---|---|---:|---:|---:|---|"]
    for c in cols:
        md.append(f"| {c['col']} | {clean(c['period'])} | {fmt(c['assets'])} | {fmt(c['el'])} | {fmt(c['gap']) if c['gap'] is not None else '-'} | {c['status']} |")
    md.append("")
    for c in cols:
        a = res["analyses"].get(c["col"])
        if c["status"] == "Gap" and a:
            md += _gap_section(res, c, a)
        elif c["status"] == "Error" and a:
            md += _error_section(res, c, a)
        elif c["status"] == "Tallies" and a and a.get("dup_only"):
            if a["items"]:
                md += [f"## {clean(c['period'])} - column {c['col']}: tallies, but a ledger is linked more than once",
                       "The totals agree, yet the checks below found the same ledger amount linked from more than one place. "
                       "A duplicate like this is normally hidden by an offsetting error somewhere else, so it is worth fixing even though the column ties.", ""]
                md += [f"- {i['text']}" for i in a["items"]] + [""]
            if a["observations"]:
                md += [f"## {clean(c['period'])} - column {c['col']}: tallies, but check these typed inputs",
                       "The balance sheet cannot reveal an error that is carried consistently into both sides (for example a wrong depreciation input that flows into the "
                       "schedule and the P&L alike). These typed numbers stand out from their peers:", ""]
                md += [f"- **{o['sev']}** - {o['text']}" for o in a["observations"]] + [""]
    obs = res["observations"]
    if obs:
        md += ["## Other things noticed (not necessarily causing a gap)"]
        for o in obs:
            md.append(f"- **{o['sev']}** - {o['text']}")
        md.append("")
    if res.get("warnings"):
        md += ["## Notes on the analysis"] + [f"- {w}" for w in res["warnings"][:10]] + [""]
    return "\n".join(md)


def _gap_section(res, c, a):
    sh = res["sheet"]
    md = [f"## {clean(c['period'])} - column {c['col']}: {_side_word(c['gap'])} by {fmt(abs(c['gap']))}",
          f"{c['tl_cell']} = {fmt(c['el'])} against {c['ta_cell']} = {fmt(c['assets'])}.", ""]
    sols = a.get("solutions") or []
    if sols:
        best = sols[0]
        kinds = Counter(i["type"] for i in best["items"])
        md.append(f"**Most likely cause.** {_cause_headline(best, kinds)}")
        md.append("")
        for i in best["items"]:
            md.append(f"- {i['text']}")
        md.append("")
        tot = sum(i["e"] for i in best["items"])
        if best["exact"]:
            md.append(f"Together these account for {fmt(abs(tot))} - the difference is explained completely (residual {fmt(abs(best['residual']), 3)}).")
        elif best["within"]:
            md.append(f"Together these account for {fmt(abs(tot))} of the {fmt(abs(c['gap']))} difference, leaving a small unexplained residual of {fmt(abs(best['residual']), 3)}. "
                      f"Residuals like this usually come from an opening balance that is typed in or carried from the previous year's statements instead of coming from the trial balance, or from rounding - compare the opening reserves/capital with the TB.")
            blob = " ".join(i["text"] for i in best["items"])
            hi = [h for h in (a.get("hardcoded_inputs") or []) if f"'{h['line']}'" in blob]
            if hi:
                md.append("")
                md.append("Typed-in (not linked) numbers feeding the same line(s) - the usual place to look for such a residual: " +
                          "; ".join(f"{h['cell']} = {fmt(h['value'])} (feeds '{h['line']}')" for h in hi[:4]) + ".")
        else:
            md.append(f"These are only a weak lead: even the closest combination leaves {fmt(abs(best['residual']), 2)} unexplained. "
                      f"The difference most likely sits in typed-in numbers or schedules that have no link to the trial balance - see the ledger tie-out below.")
        md.append("")
        alts = [s for s in sols[1:] if s["abs_residual"] <= max(3 * abs(best["abs_residual"]), res["tol"] * 10) and s["within"]]
        if alts:
            md.append("Other combinations that also fit (check which one is right):")
            for s in alts[:2]:
                md.append("- " + "; ".join(i.get("short") or i["ref"] for i in s["items"]) + f" (residual {fmt(abs(s['residual']), 3)})")
            md.append("")
    else:
        md += ["No combination of the detected anomalies explains the gap. Review the observations below and the Excel report (TB coverage).", ""]
    tie = a.get("tieout")
    if tie and (tie["verified"] or tie["unverified"]):
        md.append("**Ledger tie-out** (ledgers that no line reads, compared with the figures typed into schedules and notes):")
        for v in tie["verified"][:6]:
            md.append(f"- Agrees: {', '.join(x.strip() for x in v['rows'])} ({fmt(v['total'])}) = {v['cells']}" + ("" if v["name_confirmed"] else " (amount match only)"))
        if len(tie["verified"]) > 6:
            md.append(f"- ...and {len(tie['verified']) - 6} more groups that agree.")
        if tie["unverified"]:
            md.append("- No schedule figure reproduces: " + ", ".join(f"{u['label'].strip()} ({fmt(abs(u['net']))})" for u in tie["unverified"]) + ".")
        for c_ in tie["candidates"]:
            for alt in c_["also"][:1]:
                md.append(f"  (the same schedule amount can also be written as {alt})")
            if c_["inputs"]:
                md.append("  Typed inputs behind it: " + "; ".join(f"{i['cell']} = {fmt(i['value'])}" + (f" ({i['label']})" if i["label"] else "") for i in c_["inputs"][:6]) + ".")
        md.append("")
    mapped = a.get("tb")
    if mapped:
        md.append(f"*Coverage check on '{mapped}': {a.get('tb_rows_linked', 0)} of {a.get('tb_rows_total', 0)} ledger rows are read by this balance sheet column; "
                  f"{len(a.get('tb_unlinked', []))} are not read by any line, {len(a.get('tb_multi', []))} are read by more than one line, {len(a.get('tb_partial', []))} are only partly read.*")
        md.append("")
    extra = [o for o in a["observations"] if o["type"] not in ("ROLE",)]
    if extra:
        md.append("Other observations for this column:")
        for o in extra:
            md.append(f"- **{o['sev']}** - {o['text']}")
        md.append("")
    return md


def _cause_headline(best, kinds):
    bits = []
    names = {"DUP": "ledger(s) linked from more than one place (counted twice or cancelling out)", "UNLINKED": "ledger(s) not linked into the statements",
             "PARTIAL": "ledger(s) only partly linked", "OVERLAP": "period amount(s) counted twice",
             "INTERNAL": "total that does not add up its lines", "TIEOUT": "schedule figure(s) that do not agree with the ledger"}
    for k, n in kinds.items():
        bits.append(f"{n} {names.get(k, k)}")
    return " + ".join(bits).capitalize() + "."


def _error_section(res, c, a):
    md = [f"## {clean(c['period'])} - column {c['col']}: cannot be compared because of formula errors",
          f"{', '.join(c.get('error_cells', []))} show an error, so the totals cannot be compared. Fix the errors first, then run the check again.", ""]
    for o in a["observations"]:
        md.append(f"- **{o['sev']}** - {o['text']}")
    md.append("")
    return md


# --------------------------------------------------------------------------- Excel report
def to_excel(res) -> bytes:
    wb = Workbook()
    f_norm, f_bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="D85A30")
    f_head = Font(name="Arial", size=10, bold=True, color="FFFFFF")

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = f_head, head_fill
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = f_norm
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00;(#,##0.00);-"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    ws = wb.create_sheet("Analysis")
    ws.column_dimensions["A"].width = 150
    for line in res["markdown"].splitlines():
        ws.append([re.sub(r"\*\*|\*|^#+\s*", "", line)])
    for row in ws.iter_rows():
        for c in row:
            c.font = f_norm
            c.alignment = Alignment(wrap_text=True, vertical="top")
    sheet("Columns", ["Column", "Period", "Total Assets", "Total Equity & Liabilities", "Difference", "Status"],
          [[c["col"], clean(c["period"]), c["assets"], c["el"], c["gap"], c["status"]] for c in res["columns"]], [10, 40, 18, 22, 16, 12])
    frows = []
    for col, a in res["analyses"].items():
        for k, s in enumerate(a.get("solutions") or [], 1):
            for i in s["items"]:
                frows.append([col, f"Option {k}", i["type"], i["ref"], i["e"], i["text"], s["residual"]])
    sheet("Explanations", ["Column", "Combination", "Type", "Reference", "Effect on gap", "Description", "Residual after combination"], frows, [10, 14, 12, 28, 16, 110, 18])
    urows, mrows, prow = [], [], []
    for col, a in res["analyses"].items():
        for u in a.get("tb_unlinked", []):
            urows.append([col, a["tb"], u["row"], u["code"], u["label"], u["dr"], u["cr"], u["net"], u["hint"], u.get("typed_in", "")])
        for m in a.get("tb_multi", []):
            mrows.append([col, a["tb"], m["row"], m["label"], " | ".join(m["lines"]), " | ".join(fmt(x) for x in m["contrib"]), ", ".join(m["via"])])
        for p in a.get("tb_partial", []):
            prow.append([col, a["tb"], p["row"], p["label"], ", ".join(p["cols"]), p["net"]])
    sheet("TB rows not linked", ["Column", "TB sheet", "Row", "Code", "Ledger", "Debit", "Credit", "Net (Dr-Cr)", "Neighbours feed", "Same amount typed in at"], urows, [9, 18, 8, 10, 40, 14, 14, 14, 36, 28])
    sheet("TB rows read twice", ["Column", "TB sheet", "Row", "Ledger", "Lines that read it", "Amounts", "Through cells"], mrows, [9, 18, 8, 36, 60, 24, 50])
    sheet("TB rows part-linked", ["Column", "TB sheet", "Row", "Ledger", "Unread amount column(s)", "Net amount not read"], prow, [9, 18, 8, 36, 22, 18])
    trows = []
    for col, a in res["analyses"].items():
        t = a.get("tieout")
        if t:
            trows += [[col, "Agrees", ", ".join(x.strip() for x in v["rows"]), v["total"], v["cells"], None] for v in t["verified"]]
            trows += [[col, "Not reproduced by any schedule figure", u["label"].strip(), abs(u["net"]), "", None] for u in t["unverified"]]
            trows += [[col, "Closest schedule figure (difference equals the gap)", ", ".join(x.strip() for x in c["rows"]), c["ledger_total"], c["cells"], c["difference"]] for c in t["candidates"]]
    sheet("Ledger tie-out", ["Column", "Result", "Ledgers", "Ledger total", "Schedule cells", "Schedule minus ledger"], trows, [9, 40, 60, 16, 50, 18])
    orows = [[o["sev"], o["text"]] for o in res["observations"]]
    for col, a in res["analyses"].items():
        orows += [[o["sev"], f"[{col}] {o['text']}"] for o in a["observations"]]
        orows += [["Info", f"[{col}] {o['text']}"] for o in a.get("overlaps", [])]
    sheet("Observations", ["Severity", "Observation"], orows, [12, 150])
    hrows = [[col, h["cell"], h["value"], h["line"]] for col, a in res["analyses"].items() for h in a.get("hardcoded_inputs", [])]
    sheet("Typed-in inputs", ["Column", "Cell", "Value", "Feeds line"], hrows, [9, 30, 18, 50])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


if __name__ == "__main__":
    import sys
    path = sys.argv[1]
    sheet_arg = sys.argv[2] if len(sys.argv) > 2 else None
    bk = load_book(open(path, "rb").read(), path.split("/")[-1])
    r = analyze(bk, sheet_arg)
    print(r.get("markdown") or r.get("error"))
